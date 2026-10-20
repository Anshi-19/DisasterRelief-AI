"""Aggregate results into Excel, markdown reports, and dashboard data."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from eval_config import (
    CODEGEN_RESULTS_FILE,
    DASHBOARD_DATA_FILE,
    DASHBOARD_DIR,
    DATASET_FILE,
    EMBEDDING_MODEL,
    EVALUATION_DATE,
    EVALUATION_VERSION,
    EXPORTS_DIR,
    MODELS,
    MODEL_DISPLAY,
    NUM_PREDICT,
    PROJECT_ROOT,
    QUALITY_SCORES_FILE,
    RAG_CASE_STUDIES_FILE,
    RAW_CANONICAL_COPY,
    RAW_RESULTS_FILE,
    REPORTS_DIR,
    REPOSITORY_EVAL_FILE,
    RESOURCE_METRICS_FILE,
    RESULTS_DIR,
    RETRIEVAL_METRICS_FILE,
    SUMMARY_FILE,
    TEMPERATURE,
    TOP_K,
    V3_QUALITY_SCORES_FILE,
    V3_RAW_RESULTS_FILE,
    V3_RETRIEVAL_METRICS_FILE,
    V3_RESOURCE_METRICS_FILE,
    V3_SUMMARY_FILE,
    V3_V4_COMPARISON_FILE,
    WEEK4_DIR,
)
from measure_resources import percentile

HEADER_FILL = PatternFill("solid", fgColor="111827")
HEADER_FONT = Font(color="F9FAFB", bold=True)
ALT_FILL = PatternFill("solid", fgColor="F3F4F6")
THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def load_dataset():
    with DATASET_FILE.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def safe_mean(values):
    values = [v for v in values if v is not None]
    return round(mean(values), 4) if values else None


def fmt(value, digits=3):
    if value is None:
        return "N/A"
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return value


def style_header(ws, columns):
    for col, title in enumerate(columns, start=1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    ws.freeze_panes = "A2"
    for index in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(index)].width = min(28, max(14, len(columns[index - 1]) + 4))


def write_rows(ws, columns, rows):
    style_header(ws, columns)
    for r_index, row in enumerate(rows, start=2):
        for c_index, key in enumerate(columns, start=1):
            cell = ws.cell(r_index, c_index, row.get(key))
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r_index % 2 == 0:
                cell.fill = ALT_FILL
    ws.row_dimensions[1].height = 28


def factual_labels(model_stats):
    labels = {model: [] for model in model_stats}
    def pick(metric, higher=True, label=None):
        usable = {m: v[metric] for m, v in model_stats.items() if v.get(metric) is not None}
        if not usable:
            return
        best = max(usable, key=usable.get) if higher else min(usable, key=usable.get)
        # Require uniqueness
        winners = [m for m, val in usable.items() if val == usable[best]]
        if len(winners) == 1:
            labels[best].append(label)

    pick("correctness", True, "Highest measured correctness")
    pick("relevance", True, "Highest measured relevance")
    pick("hit_at_k", True, "Highest retrieval Hit@K")
    pick("mean_latency", False, "Lowest measured mean latency")
    pick("median_latency", False, "Lowest measured median latency")
    pick("test_pass_rate", True, "Highest code test-pass rate")
    pick("throughput", True, "Highest measured generation throughput")
    pick("failure_rate", False, "Lowest measured failure rate")
    return labels


def build_summary(raw, quality, retrieval, resources, repo, codegen):
    q_models = quality["aggregates"]["models"]
    r_models = retrieval["aggregates"]["models"]
    res_models = resources["models"]
    cg = quality["aggregates"].get("code_generation", {})
    successful = sum(1 for item in raw if item.get("success"))
    failed = len(raw) - successful
    model_stats = {}
    for model in MODELS:
        q = q_models.get(model, {})
        r = r_models.get(model, {})
        p = res_models.get(model, {})
        c = cg.get(model, {})
        model_stats[model] = {
            "display": MODEL_DISPLAY.get(model, model),
            "correctness": q.get("correctness"),
            "relevance": q.get("relevance"),
            "hallucination": q.get("hallucination"),
            "hallucination_rate_any": q.get("hallucination_rate_any"),
            "hit_at_k": r.get("hit_at_k"),
            "mrr": r.get("mrr"),
            "precision_at_k": r.get("precision_at_k"),
            "recall_at_k": r.get("recall_at_k"),
            "mean_latency": p.get("average_latency"),
            "median_latency": p.get("median_latency"),
            "p95_latency": p.get("p95_latency"),
            "retrieval_latency": p.get("average_retrieval_latency"),
            "prompt_tokens": p.get("average_prompt_tokens"),
            "generated_tokens": p.get("average_generated_tokens"),
            "total_tokens": p.get("average_total_tokens"),
            "throughput": p.get("average_throughput"),
            "memory_mb": p.get("average_memory_mb"),
            "cpu_percent": p.get("average_cpu_percent"),
            "failure_rate": p.get("failure_rate"),
            "successful": p.get("successful"),
            "failed": p.get("failed"),
            "experiments": p.get("experiments"),
            "test_pass_rate": c.get("test_pass_rate"),
            "task_logic_correctness": c.get("mean_task_logic_correctness"),
            "load_duration": p.get("average_model_load_duration"),
            "prompt_eval_duration": p.get("average_prompt_eval_duration"),
        }
    category_matrix = []
    categories = sorted({item["category"] for item in raw})
    lat_index = defaultdict(list)
    for item in raw:
        if item.get("success") and item.get("response_latency") is not None:
            lat_index[(item["model"], item["category"])].append(item["response_latency"])
    for category in categories:
        for model in MODELS:
            key = f"{model}::{category}"
            q = quality["aggregates"]["by_category"].get(key, {})
            r = retrieval["aggregates"]["by_category"].get(key, {})
            category_matrix.append(
                {
                    "category": category,
                    "model": model,
                    "display": MODEL_DISPLAY.get(model, model),
                    "correctness": q.get("correctness"),
                    "relevance": q.get("relevance"),
                    "hallucination": q.get("hallucination"),
                    "hit_at_k": r.get("hit_at_k"),
                    "mrr": r.get("mrr"),
                    "latency": round(mean(lat_index[(model, category)]), 3)
                    if lat_index[(model, category)]
                    else None,
                }
            )
    return {
        "evaluation_version": EVALUATION_VERSION,
        "evaluation_date": EVALUATION_DATE,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "questions": 28,
        "models": MODELS,
        "planned_experiments": 84,
        "recorded_experiments": len(raw),
        "successful_experiments": successful,
        "failed_experiments": failed,
        "embedding_model": EMBEDDING_MODEL,
        "top_k": TOP_K,
        "temperature": TEMPERATURE,
        "num_predict": NUM_PREDICT,
        "model_stats": model_stats,
        "factual_labels": factual_labels(model_stats),
        "category_matrix": category_matrix,
        "gpu_note": resources.get("gpu_note"),
        "resource_method": resources.get("measurement_method"),
        "repository_aggregates": repo.get("aggregates", {}),
        "score_definitions": quality["aggregates"]["score_definitions"],
    }


def _safe_load(path: Path):
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


def build_v3_v4_comparison(v4_summary):
    v3_quality = _safe_load(V3_QUALITY_SCORES_FILE) or {}
    v3_retrieval = _safe_load(V3_RETRIEVAL_METRICS_FILE) or {}
    v3_resources = _safe_load(V3_RESOURCE_METRICS_FILE) or {}
    v3_summary = _safe_load(V3_SUMMARY_FILE) or {}
    v3_models = (v3_quality.get("aggregates") or {}).get("models") or {}
    v3_retrieval_models = (v3_retrieval.get("aggregates") or {}).get("models") or {}
    v3_resource_models = v3_resources.get("models") or {}
    v3_codegen = (v3_quality.get("aggregates") or {}).get("code_generation") or {}
    v3_categories = (v3_quality.get("aggregates") or {}).get("by_category") or {}
    comparison = {
        "v3_version": "week4_v3_controlled",
        "v4_version": EVALUATION_VERSION,
        "models": {},
        "categories": [],
        "notes": [
            "V3 raw results were not modified.",
            "V4 uses the same 28 questions, ground truth, and scoring scale.",
            "A documented evaluator bug was fixed: unittest.main() OK/exit 0 is now counted as a code-test pass.",
        ],
    }
    for model in MODELS:
        v3q = v3_models.get(model, {})
        v3r = v3_retrieval_models.get(model, {})
        v3p = v3_resource_models.get(model, {})
        v3c = v3_codegen.get(model, {})
        v4s = v4_summary["model_stats"].get(model, {})
        comparison["models"][model] = {
            "display": MODEL_DISPLAY.get(model, model),
            "v3": {
                "correctness": v3q.get("correctness"),
                "relevance": v3q.get("relevance"),
                "hit_at_k": v3r.get("hit_at_k"),
                "mrr": v3r.get("mrr"),
                "hallucination": v3q.get("hallucination"),
                "test_pass_rate": v3c.get("test_pass_rate"),
                "task_logic_correctness": v3c.get("mean_task_logic_correctness"),
                "mean_latency": v3p.get("average_latency"),
                "median_latency": v3p.get("median_latency"),
                "p95_latency": v3p.get("p95_latency"),
                "throughput": v3p.get("average_throughput"),
                "memory_mb": v3p.get("average_memory_mb"),
                "cpu_percent": v3p.get("average_cpu_percent"),
            },
            "v4": {
                "correctness": v4s.get("correctness"),
                "relevance": v4s.get("relevance"),
                "hit_at_k": v4s.get("hit_at_k"),
                "mrr": v4s.get("mrr"),
                "hallucination": v4s.get("hallucination"),
                "test_pass_rate": v4s.get("test_pass_rate"),
                "task_logic_correctness": v4s.get("task_logic_correctness"),
                "mean_latency": v4s.get("mean_latency"),
                "median_latency": v4s.get("median_latency"),
                "p95_latency": v4s.get("p95_latency"),
                "throughput": v4s.get("throughput"),
                "memory_mb": v4s.get("memory_mb"),
                "cpu_percent": v4s.get("cpu_percent"),
            },
        }
    categories = sorted({row["category"] for row in v4_summary.get("category_matrix", [])})
    for category in categories:
        entry = {"category": category, "models": {}}
        for model in MODELS:
            v3_item = v3_categories.get(f"{model}::{category}", {})
            v4_item = next(
                (
                    row
                    for row in v4_summary.get("category_matrix", [])
                    if row["model"] == model and row["category"] == category
                ),
                {},
            )
            entry["models"][model] = {
                "v3_correctness": v3_item.get("correctness"),
                "v4_correctness": v4_item.get("correctness"),
            }
        comparison["categories"].append(entry)
    comparison["v3_recorded_experiments"] = v3_summary.get("recorded_experiments")
    comparison["v4_recorded_experiments"] = v4_summary.get("recorded_experiments")
    return comparison


def write_excel(dataset, raw, quality, retrieval, resources, rag, repo, summary):
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Executive Summary"
    rows = [
        ["Field", "Value"],
        ["Project", "DisasterRelief-AI"],
        ["Evaluation version", EVALUATION_VERSION],
        ["Evaluation date", EVALUATION_DATE],
        ["Questions", 28],
        ["Models", ", ".join(MODELS)],
        ["Planned experiments", 84],
        ["Recorded experiments", summary["recorded_experiments"]],
        ["Successful experiments", summary["successful_experiments"]],
        ["Failed experiments", summary["failed_experiments"]],
        ["Embedding model", EMBEDDING_MODEL],
        ["top_k", TOP_K],
        ["temperature", TEMPERATURE],
        ["num_predict", NUM_PREDICT],
        ["GPU", summary["gpu_note"]],
    ]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            cell.border = THIN
            if r == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
    ws["A17"] = "Model"
    ws["B17"] = "Correctness"
    ws["C17"] = "Relevance"
    ws["D17"] = "Hit@K"
    ws["E17"] = "Hallucination"
    ws["F17"] = "Mean latency"
    ws["G17"] = "Median latency"
    ws["H17"] = "p95 latency"
    ws["I17"] = "Failure rate"
    for col in range(1, 10):
        cell = ws.cell(17, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for index, model in enumerate(MODELS, start=18):
        stats = summary["model_stats"][model]
        ws.cell(index, 1, model)
        ws.cell(index, 2, stats["correctness"])
        ws.cell(index, 3, stats["relevance"])
        ws.cell(index, 4, stats["hit_at_k"])
        ws.cell(index, 5, stats["hallucination"])
        ws.cell(index, 6, stats["mean_latency"])
        ws.cell(index, 7, stats["median_latency"])
        ws.cell(index, 8, stats["p95_latency"])
        ws.cell(index, 9, stats["failure_rate"])

    ds = wb.create_sheet("Dataset")
    ds_cols = list(dataset[0].keys()) if dataset else []
    write_rows(ds, ds_cols, dataset)

    raw_sheet = wb.create_sheet("Raw Model Results")
    raw_cols = [
        "model",
        "question_id",
        "category",
        "success",
        "error",
        "answer",
        "retrieval_source",
        "retrieval_latency",
        "response_latency",
        "prompt_tokens",
        "generated_tokens",
        "total_tokens",
        "throughput",
        "generation_duration",
        "model_load_duration",
        "prompt_eval_duration",
        "retrieved_chunk_count",
        "attempts",
        "http_status",
    ]
    raw_rows = [{key: item.get(key) for key in raw_cols} for item in raw]
    write_rows(raw_sheet, raw_cols, raw_rows)
    last_raw = 1 + len(raw_rows)

    q_sheet = wb.create_sheet("Quality Metrics")
    q_sheet["A1"] = "model"
    q_sheet["B1"] = "correctness_avg"
    q_sheet["C1"] = "relevance_avg"
    q_sheet["D1"] = "hallucination_avg"
    for col in range(1, 5):
        q_sheet.cell(1, col).fill = HEADER_FILL
        q_sheet.cell(1, col).font = HEADER_FONT
    per_q = quality["per_question"]
    detail = wb.create_sheet("Quality Detail")
    detail_cols = [
        "model",
        "question_id",
        "category",
        "correctness",
        "relevance",
        "hallucination",
        "correctness_method",
        "relevance_method",
        "hallucination_method",
        "correctness_rationale",
        "success",
    ]
    write_rows(detail, detail_cols, per_q)
    last_q = 1 + len(per_q)
    for index, model in enumerate(MODELS, start=2):
        q_sheet.cell(index, 1, model)
        q_sheet.cell(index, 2, f"=AVERAGEIF('Quality Detail'!A:A,A{index},'Quality Detail'!D:D)")
        q_sheet.cell(index, 3, f"=AVERAGEIF('Quality Detail'!A:A,A{index},'Quality Detail'!E:E)")
        q_sheet.cell(index, 4, f"=AVERAGEIF('Quality Detail'!A:A,A{index},'Quality Detail'!F:F)")

    ret_sheet = wb.create_sheet("Retrieval Metrics")
    ret_cols = [
        "model",
        "question_id",
        "category",
        "hit_at_k",
        "first_relevant_rank",
        "reciprocal_rank",
        "precision_at_k",
        "recall_at_k",
        "expected_files",
        "retrieved_files",
        "matched_expected_files",
    ]
    ret_rows = []
    for item in retrieval["per_question"]:
        row = dict(item)
        row["expected_files"] = "; ".join(item.get("expected_files") or [])
        row["retrieved_files"] = "; ".join(item.get("retrieved_files") or [])
        row["matched_expected_files"] = "; ".join(item.get("matched_expected_files") or [])
        ret_rows.append(row)
    write_rows(ret_sheet, ret_cols, ret_rows)

    perf = wb.create_sheet("Performance Metrics")
    write_rows(
        perf,
        [
            "model",
            "experiments",
            "successful",
            "failed",
            "failure_rate",
            "average_latency",
            "median_latency",
            "p95_latency",
            "average_retrieval_latency",
            "average_prompt_tokens",
            "average_generated_tokens",
            "average_total_tokens",
            "average_throughput",
            "average_model_load_duration",
            "average_prompt_eval_duration",
        ],
        [{"model": model, **values} for model, values in resources["models"].items()],
    )
    perf["A20"] = "Latency formulas over successful Raw Model Results rows"
    perf["A21"] = "Model"
    perf["B21"] = "AVERAGE latency"
    perf["C21"] = "Note"
    for index, model in enumerate(MODELS, start=22):
        perf.cell(index, 1, model)
        perf.cell(
            index,
            2,
            f"=AVERAGEIF('Raw Model Results'!A:A,A{index},'Raw Model Results'!I:I)",
        )
        perf.cell(index, 3, "Calculated from Raw Model Results.response_latency")

    res_sheet = wb.create_sheet("Resource Metrics")
    res_sheet["A1"] = "Measurement method"
    res_sheet["A2"] = resources.get("measurement_method")
    res_sheet["A4"] = "GPU note"
    res_sheet["A5"] = resources.get("gpu_note")
    write_rows(
        wb.create_sheet("Category Analysis"),
        [
            "category",
            "model",
            "display",
            "correctness",
            "relevance",
            "hallucination",
            "hit_at_k",
            "mrr",
            "latency",
        ],
        summary["category_matrix"],
    )

    cg_rows = []
    if CODEGEN_RESULTS_FILE.exists():
        cg_rows = load_json(CODEGEN_RESULTS_FILE)
        for row in cg_rows:
            row["logic_notes"] = "; ".join(row.get("logic_notes") or [])
    write_rows(
        wb.create_sheet("Code Generation"),
        [
            "model",
            "question_id",
            "syntax_valid",
            "executed",
            "passed",
            "error_type",
            "task_logic_correctness",
            "logic_notes",
            "error",
        ],
        cg_rows,
    )

    rag_rows = []
    for case in rag.get("cases", []):
        rag_rows.append(
            {
                "question_id": case.get("question_id"),
                "model": case.get("model"),
                "behavior": case.get("behavior"),
                "question": case.get("question"),
                "ground_truth": case.get("ground_truth"),
                "model_response": case.get("model_response"),
                "key_relevant_information": "; ".join(case.get("key_relevant_information") or []),
                "missed_information": "; ".join(case.get("missed_information") or []),
                "correctness": (case.get("scores") or {}).get("correctness"),
                "relevance": (case.get("scores") or {}).get("relevance"),
                "hallucination": (case.get("scores") or {}).get("hallucination"),
                "hit_at_k": (case.get("retrieval") or {}).get("hit_at_k"),
                "analysis": case.get("analysis"),
            }
        )
    write_rows(
        wb.create_sheet("RAG Case Studies"),
        list(rag_rows[0].keys()) if rag_rows else ["question_id"],
        rag_rows,
    )

    write_rows(
        wb.create_sheet("Repository Evaluation"),
        [
            "model",
            "question_id",
            "file_identification_score",
            "dependency_reasoning_score",
            "overall_score",
            "expected_files",
            "identified_files_in_answer",
            "success",
        ],
        [
            {
                **item,
                "expected_files": "; ".join(item.get("expected_files") or []),
                "identified_files_in_answer": "; ".join(item.get("identified_files_in_answer") or []),
            }
            for item in repo.get("per_question", [])
        ],
    )

    comparison = summary.get("v3_v4_comparison")
    if comparison:
        compare_rows = []
        for model, payload in comparison.get("models", {}).items():
            v3 = payload.get("v3") or {}
            v4 = payload.get("v4") or {}
            compare_rows.append(
                {
                    "model": model,
                    "display": payload.get("display"),
                    "v3_correctness": v3.get("correctness"),
                    "v4_correctness": v4.get("correctness"),
                    "v3_relevance": v3.get("relevance"),
                    "v4_relevance": v4.get("relevance"),
                    "v3_hit_at_k": v3.get("hit_at_k"),
                    "v4_hit_at_k": v4.get("hit_at_k"),
                    "v3_mrr": v3.get("mrr"),
                    "v4_mrr": v4.get("mrr"),
                    "v3_hallucination": v3.get("hallucination"),
                    "v4_hallucination": v4.get("hallucination"),
                    "v3_code_pass": v3.get("test_pass_rate"),
                    "v4_code_pass": v4.get("test_pass_rate"),
                    "v3_task_logic": v3.get("task_logic_correctness"),
                    "v4_task_logic": v4.get("task_logic_correctness"),
                    "v3_mean_latency": v3.get("mean_latency"),
                    "v4_mean_latency": v4.get("mean_latency"),
                    "v3_median_latency": v3.get("median_latency"),
                    "v4_median_latency": v4.get("median_latency"),
                    "v3_p95_latency": v3.get("p95_latency"),
                    "v4_p95_latency": v4.get("p95_latency"),
                    "v3_throughput": v3.get("throughput"),
                    "v4_throughput": v4.get("throughput"),
                    "v3_memory": v3.get("memory_mb"),
                    "v4_memory": v4.get("memory_mb"),
                    "v3_cpu": v3.get("cpu_percent"),
                    "v4_cpu": v4.get("cpu_percent"),
                }
            )
        write_rows(
            wb.create_sheet("V3 vs V4"),
            list(compare_rows[0].keys()) if compare_rows else ["model"],
            compare_rows,
        )
        cat_rows = []
        for entry in comparison.get("categories", []):
            for model, values in entry.get("models", {}).items():
                cat_rows.append(
                    {
                        "category": entry.get("category"),
                        "model": model,
                        "v3_correctness": values.get("v3_correctness"),
                        "v4_correctness": values.get("v4_correctness"),
                    }
                )
        write_rows(
            wb.create_sheet("V3 vs V4 Categories"),
            ["category", "model", "v3_correctness", "v4_correctness"],
            cat_rows,
        )

    method = wb.create_sheet("Methodology")
    method["A1"] = "Item"
    method["B1"] = "Setting"
    method_rows = [
        ("Same questions for all models", "Yes, 28 frozen dataset items"),
        ("Retrieval configuration", f"top_k={TOP_K}, embedding={EMBEDDING_MODEL}"),
        ("Repository collection", "disaster_relief_code"),
        ("Disaster collection", "disaster_relief (Week 3, unmodified)"),
        ("Temperature", TEMPERATURE),
        ("Max output tokens", NUM_PREDICT),
        ("Retries", "up to 2 retries on transient Ollama errors"),
        ("Correctness scale", "1.0 / 0.5 / 0.0 rubric"),
        ("Hallucination vs context-use", "Wrong selection from retrieved facts is not labeled hallucination"),
        ("Code tests", "Extracted Python executed against backend/ with PYTHONPATH set"),
        ("GPU", summary["gpu_note"]),
        ("Quality averages", "Failed generations excluded"),
        ("Excel formulas", "Quality Metrics!B2:D4 and Performance AVERAGEIF over detail/raw sheets"),
        ("Raw last row", last_raw),
        ("Quality detail last row", last_q),
    ]
    method["A1"].fill = HEADER_FILL
    method["B1"].fill = HEADER_FILL
    method["A1"].font = HEADER_FONT
    method["B1"].font = HEADER_FONT
    for index, (item, value) in enumerate(method_rows, start=2):
        method.cell(index, 1, item)
        method.cell(index, 2, value)
        method.column_dimensions["A"].width = 40
        method.column_dimensions["B"].width = 80

    path = EXPORTS_DIR / "week4_evaluation.xlsx"
    wb.save(path)
    return path


def category_interpretation(matrix):
    by_cat = defaultdict(list)
    for row in matrix:
        by_cat[row["category"]].append(row)
    text = {}
    for category, rows in by_cat.items():
        best = max(rows, key=lambda item: item["correctness"] or -1)
        slowest = max(rows, key=lambda item: item["latency"] or -1)
        text[category] = (
            f"In this benchmark, {MODEL_DISPLAY.get(best['model'], best['model'])} "
            f"recorded the highest rubric correctness for {category} "
            f"({fmt(best['correctness'])}). "
            f"{MODEL_DISPLAY.get(slowest['model'], slowest['model'])} recorded the "
            f"highest mean latency in this category ({fmt(slowest['latency'], 2)} s)."
        )
    return text


def write_reports(summary, quality, retrieval, resources, rag, repo):
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    interpretations = category_interpretation(summary["category_matrix"])
    lines = [
        "# DisasterRelief-AI Week 4 Technical Report",
        "",
        "## 1. Introduction",
        "",
        "This report documents a controlled Week 4 evaluation of local Ollama code models on the existing DisasterRelief-AI Week 3 RAG application. All quantitative figures below are computed from `week4_evaluation/results/` after actually running retrieval, generation, scoring, and resource sampling. No scores were edited to favor a model.",
        "",
        "## 2. Week 3 System Overview",
        "",
        "The live application path is:",
        "",
        "1. `frontend/src/App.jsx` `analyzeSituation()` posts `{message}` to `POST /api/chat`.",
        "2. `backend/main.py` `chat()` calls `process_question(request.message, top_k=8)`.",
        "3. `process_question()` in `orchestration_service.py` calls `retrieve_documents()`.",
        "4. If retrieval is empty, it returns the insufficient-knowledge fallback and does not call the LLM.",
        "5. Otherwise `build_context()` formats document/authority/disaster/pages/content and `generate_response()` posts to Ollama.",
        "6. Retrieval uses `nomic-embed-text:latest`, Chroma collection `disaster_relief`, optional `detect_disaster()` metadata filter, distance cutoff 350, and per-document limits `max(2, top_k // 2)`.",
        "7. The application LLM remains `codellama:latest`. Week 4 benchmarking did not permanently replace that setting.",
        "",
        "`backend/rag/` contains an older RAG implementation. It was indexed for repository search but is not the live `/api/chat` path.",
        "",
        "## 3. Week 4 Objectives",
        "",
        "- Evaluate the same 28 tasks on at least three local models.",
        "- Measure quality, retrieval, latency, tokens, and host resources.",
        "- Analyze RAG traces and multi-file repository understanding.",
        "- Produce reproducible JSON, Excel, markdown, and dashboard artifacts.",
        "",
        "## 4. System Architecture",
        "",
        "Live Week 3 path (used by the application):",
        "",
        "`frontend/src/App.jsx` `analyzeSituation()` → `POST /api/chat` in `backend/main.py` → `process_question()` → `retrieve_documents()` (Chroma `disaster_relief`) → `build_context()` → `generate_response()` → Ollama `codellama:latest`.",
        "",
        "Legacy path (indexed, not called by `/api/chat`): `backend/rag/rag_service.py`, `backend/rag/retrieval.py`, `backend/rag/embeddings.py`.",
        "",
        "Week 4 evaluation retrieves repository questions from collection `disaster_relief_code` and RAG questions through the live `retrieve_documents()` function. Retrieval is executed once per question; all three models receive the same frozen context.",
        "",
        "## 5. Evaluation Environment",
        "",
        f"- Date: {EVALUATION_DATE}",
        f"- Version: `{EVALUATION_VERSION}`",
        f"- Host: local macOS with Ollama",
        f"- Embedding model: `{EMBEDDING_MODEL}`",
        f"- Retrieval top-k: {TOP_K}",
        f"- Temperature: {TEMPERATURE}",
        f"- `num_predict`: {NUM_PREDICT}",
        f"- Disaster collection: `vector_store/chroma_db` / `disaster_relief`",
        f"- Code collection: `vector_store` / `disaster_relief_code`",
        "",
        f"{summary['gpu_note']}",
        "",
        f"Resource method: {summary['resource_method']}",
        "",
        "## 6. Models",
        "",
    ]
    for model in MODELS:
        lines.append(f"- `{model}` (display: {MODEL_DISPLAY.get(model, model)})")
    lines += [
        "",
        "Fairness: identical questions, prompt wrapper, top-k, temperature, max tokens, embedding model, and frozen retrieval context per question.",
        "",
        "## 7. Dataset",
        "",
        "28 tasks, 7 categories × 4 items: E01–E04, CR01–CR04, D01–D04, B01–B04, CG01–CG04, RF01–RF04, RG01–RG04. Each row stores ground truth, relevant files, and explicit correctness/relevance/hallucination/retrieval rubrics. Code-generation items specify required imports, function use, inputs, expected values, and syntactic validity.",
        "",
        "## 8. Experimental Methodology",
        "",
        "Repository questions retrieve from `disaster_relief_code`. RAG questions retrieve through the Week 3 `retrieve_documents()` function. Each Ollama `generate` call is retried up to two times on transient errors. Persistent failures are stored with HTTP/error text, counted in failure rate, and excluded from quality averages.",
        "",
        "## 9. Metrics and Definitions",
        "",
        json.dumps(summary["score_definitions"], indent=2),
        "",
        "Retrieval Hit@K, MRR, Precision@K and Recall@K use dataset `relevant_files`, not answer keyword overlap.",
        "",
        "## 10. Retrieval Methodology",
        "",
        "Hit@K is 1 if at least one ground-truth relevant file or document appears in the top-K retrieved chunks. MRR is the mean of 1/rank of the first relevant result (0 if none). Precision@K is the fraction of retrieved chunks whose source matches a ground-truth item. Recall@K is the fraction of ground-truth items found in the retrieved set.",
        "",
        "Because retrieval is executed once per question and the same chunks are passed to every model, Hit@K, MRR, Precision@K, and Recall@K are identical across models for the same question. Model differences therefore come from generation, not from different retrieval.",
        "",
        "The live application collection `disaster_relief` is stored under `vector_store/chroma_db`. The Week 4 code collection `disaster_relief_code` is stored under `vector_store` and does not replace the disaster knowledge base.",
        "",
        "## 11. Quantitative Results",
        "",
        f"- Recorded experiments: {summary['recorded_experiments']}",
        f"- Successful: {summary['successful_experiments']}",
        f"- Failed: {summary['failed_experiments']}",
        "",
        "| Model | Correctness | Relevance | Hit@K | MRR | Hallucination | Code pass | Mean lat. | Median | p95 | Tok/s | Fail |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for model in MODELS:
        s = summary["model_stats"][model]
        lines.append(
            "| {display} | {c} | {r} | {h} | {mrr} | {hall} | {cg} | {lat} | {med} | {p95} | {tps} | {fail} |".format(
                display=s["display"],
                c=fmt(s["correctness"]),
                r=fmt(s["relevance"]),
                h=fmt(s["hit_at_k"]),
                mrr=fmt(s["mrr"]),
                hall=fmt(s["hallucination"]),
                cg=fmt(s["test_pass_rate"]),
                lat=fmt(s["mean_latency"], 2),
                med=fmt(s["median_latency"], 2),
                p95=fmt(s["p95_latency"], 2),
                tps=fmt(s["throughput"], 2),
                fail=fmt(s["failure_rate"], 4),
            )
        )
    lines += [
        "",
        "These values are **measured or rubric-derived**. Interpretation follows in later sections and does not change the raw JSON.",
        "",
        "## 12. Category-wise Analysis",
        "",
        "| Category | Model | Correctness | Relevance | Retrieval Hit@K | Hallucination | Latency |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for row in summary["category_matrix"]:
        lines.append(
            f"| {row['category']} | {row['display']} | {fmt(row['correctness'])} | {fmt(row['relevance'])} | {fmt(row['hit_at_k'])} | {fmt(row['hallucination'])} | {fmt(row['latency'], 2)} |"
        )
    lines.append("")
    for category, text in interpretations.items():
        lines.append(f"**{category}.** {text}")
        lines.append("")
    lines += [
        "## 13. Code Generation Analysis",
        "",
        "CG01–CG04 were executed against the real `backend/` package. Import-path failures are reported separately from task-logic correctness.",
        "",
    ]
    for model, values in quality["aggregates"].get("code_generation", {}).items():
        lines.append(
            f"- `{model}`: test_pass_rate={fmt(values.get('test_pass_rate'))}, "
            f"functional_pass_rate={fmt(values.get('functional_pass_rate'))}, "
            f"mean_task_logic_correctness={fmt(values.get('mean_task_logic_correctness'))}, "
            f"passed={values.get('passed')}/{values.get('n')}, errors={values.get('error_types')}"
        )
    lines += [
        "",
        "## 14. RAG Pipeline Analysis",
        "",
        "Inspector cases are stored in `results/rag_case_studies.json` and were selected from actual traces.",
        "",
    ]
    for case in rag.get("cases", []):
        scores = case.get("scores") or {}
        lines += [
            f"### {case.get('question_id')} / {case.get('model')} — {case.get('behavior')}",
            "",
            f"**Question.** {case.get('question')}",
            "",
            f"**Ground truth.** {case.get('ground_truth')}",
            "",
            f"**Key relevant facts in context.** {', '.join(case.get('key_relevant_information') or []) or 'None matched'}",
            "",
            f"**Missed facts.** {', '.join(case.get('missed_information') or []) or 'None'}",
            "",
            f"**Model response.** {case.get('model_response')}",
            "",
            f"**Scores.** correctness={fmt(scores.get('correctness'))}, relevance={fmt(scores.get('relevance'))}, hallucination={fmt(scores.get('hallucination'))}, hit@K={(case.get('retrieval') or {}).get('hit_at_k')}",
            "",
            f"**Analysis.** {case.get('analysis')}",
            "",
        ]
    lines += [
        "## 15. Repository-level Analysis",
        "",
        "Multi-file items evaluate whether models identify `App.jsx` → `/api/chat` → `process_question` → `retrieve_documents` → `build_context` → `generate_response` → Ollama, plus disaster detection, Chroma access, `top_k`, fallback, and embedding dependencies.",
        "",
    ]
    for model, values in repo.get("aggregates", {}).items():
        lines.append(
            f"- `{model}`: file identification={fmt(values.get('mean_file_identification_score'))}, "
            f"dependency reasoning={fmt(values.get('mean_dependency_reasoning_score'))}, "
            f"overall={fmt(values.get('mean_overall_score'))}"
        )
    lines += [
        "",
        "## 16. Quality–Latency–Resource Trade-off",
        "",
        "The scatter comparison in the dashboard plots measured mean latency against rubric correctness for this benchmark only. It is not a universal ranking.",
        "",
    ]
    for model in MODELS:
        s = summary["model_stats"][model]
        labels = summary["factual_labels"].get(model) or []
        label_text = "; ".join(labels) if labels else "no unique extremal label in this benchmark"
        lines.append(
            f"- {s['display']}: correctness={fmt(s['correctness'])}, mean latency={fmt(s['mean_latency'], 2)} s, "
            f"median={fmt(s['median_latency'], 2)} s, p95={fmt(s['p95_latency'], 2)} s, "
            f"memory={fmt(s['memory_mb'], 1)} MB, CPU={fmt(s['cpu_percent'], 1)}%, "
            f"throughput={fmt(s['throughput'], 2)} tok/s. Labels: {label_text}."
        )
    lines += [
        "",
        "## 17. Failure Analysis",
        "",
        f"- Generation failures in the `{EVALUATION_VERSION}` run: {summary['failed_experiments']} of {summary['recorded_experiments']}.",
        "- Code-generation scoring now treats a unittest that reports OK with exit code 0 as a pass, even if `unittest.main()` exited before the harness printed `TEST_PASS`. This is a documented evaluator bugfix, not a rubric change.",
        "- `week4_v3_controlled` remains historical and was not overwritten.",
        "",
        "## 18. Limitations",
        "",
        "- Automatic rubric scoring can under-credit a verbose but correct answer that omits a required token, and can over-credit keyword overlap. A human-review template is provided.",
        f"- `num_predict={NUM_PREDICT}` is shared across all three models.",
        "- CPU percent is sampled from Ollama processes and can exceed 100 on multi-core hosts.",
        "- GPU/VRAM is not reported.",
        "- Nested Python methods are indexed inside class chunks, which can reduce symbol-level recall.",
        "- Old `backend/rag/` files can appear in repository retrieval even though they are not on the live API path.",
        "",
        "## 19. V3 vs V4",
        "",
        "Historical `week4_v3_controlled` is unchanged. The numbers below compare that frozen run with this `week4_v4_improved` run.",
        "",
        "| Model | V3 corr. | V4 corr. | V3 rel. | V4 rel. | V3 Hit@K | V4 Hit@K | V3 hall. | V4 hall. | V3 code pass | V4 code pass |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    comparison = summary.get("v3_v4_comparison") or {}
    for model in MODELS:
        payload = (comparison.get("models") or {}).get(model, {})
        v3 = payload.get("v3") or {}
        v4 = payload.get("v4") or {}
        lines.append(
            "| {display} | {c3} | {c4} | {r3} | {r4} | {h3} | {h4} | {ha3} | {ha4} | {p3} | {p4} |".format(
                display=MODEL_DISPLAY.get(model, model),
                c3=fmt(v3.get("correctness")),
                c4=fmt(v4.get("correctness")),
                r3=fmt(v3.get("relevance")),
                r4=fmt(v4.get("relevance")),
                h3=fmt(v3.get("hit_at_k")),
                h4=fmt(v4.get("hit_at_k")),
                ha3=fmt(v3.get("hallucination")),
                ha4=fmt(v4.get("hallucination")),
                p3=fmt(v3.get("test_pass_rate")),
                p4=fmt(v4.get("test_pass_rate")),
            )
        )
    lines += [
        "",
        "| Category | Model | V3 correctness | V4 correctness |",
        "| --- | --- | --- | --- |",
    ]
    for entry in comparison.get("categories") or []:
        for model in MODELS:
            values = (entry.get("models") or {}).get(model, {})
            lines.append(
                f"| {entry.get('category')} | {MODEL_DISPLAY.get(model, model)} | "
                f"{fmt(values.get('v3_correctness'))} | {fmt(values.get('v4_correctness'))} |"
            )
    lines += [
        "",
        "## 20. Conclusions",
        "",
        "Conclusions are restricted to this 28-task local benchmark:",
        "",
    ]
    for model in MODELS:
        labels = summary["factual_labels"].get(model) or []
        if labels:
            lines.append(f"- `{model}`: " + "; ".join(labels) + ".")
    lines.append("")
    lines.append("No model is described as universally best.")
    lines += [
        "",
        "## 21. Reproducibility",
        "",
        "From the repository root, with Ollama running. Do not rerun `create_dataset.py`; the 28 questions are frozen.",
        "",
        "```bash",
        "source backend/venv/bin/activate",
        "python week4_evaluation/scripts/index_repository.py",
        "python week4_evaluation/scripts/test_repository_retrieval.py",
        "python week4_evaluation/scripts/run_evaluation.py",
        "python week4_evaluation/scripts/evaluate_retrieval.py",
        "python week4_evaluation/scripts/evaluate_quality.py",
        "python week4_evaluation/scripts/measure_resources.py",
        "python week4_evaluation/scripts/analyze_rag.py",
        "python week4_evaluation/scripts/analyze_repository.py",
        "python week4_evaluation/scripts/generate_report.py",
        "```",
        "",
        "Open `week4_evaluation/dashboard/index.html` in a browser.",
    ]
    report_path = REPORTS_DIR / "week4_v4_report.md"
    report_path.write_text("\n".join(lines), encoding="utf-8")

    summary_lines = [
        "# Week 4 V4 Evaluation Summary",
        "",
        f"**Project:** DisasterRelief-AI",
        f"**Objective:** Empirically evaluate three local Ollama code models on the existing Week 3 RAG application using a frozen 28-task dataset.",
        f"**AUTHORITATIVE VERSION:** `{EVALUATION_VERSION}` ({EVALUATION_DATE}). `week4_v3_controlled` is historical and was not modified.",
        "",
        "- Evaluation tasks: 28",
        "- Models: 3 (`codellama:latest`, `qwen2.5-coder:7b`, `deepseek-coder:6.7b`)",
        "- Categories: 7",
        f"- Planned experiments: 84",
        f"- Recorded experiments: {summary['recorded_experiments']}",
        f"- Successful experiments: {summary['successful_experiments']}",
        f"- Failed experiments: {summary['failed_experiments']}",
        f"- Embedding: `{EMBEDDING_MODEL}`; top-k={TOP_K}; temperature={TEMPERATURE}; num_predict={NUM_PREDICT}",
        "",
        "Key measured findings (this benchmark only):",
        "",
    ]
    for model in MODELS:
        s = summary["model_stats"][model]
        summary_lines.append(
            f"- {s['display']}: correctness={fmt(s['correctness'])}, "
            f"Hit@K={fmt(s['hit_at_k'])}, mean latency={fmt(s['mean_latency'], 2)} s, "
            f"code pass={fmt(s['test_pass_rate'])}, fail={fmt(s['failure_rate'], 4)}"
        )
    summary_lines += [
        "",
        "Retrieval Hit@K is identical across models because retrieval is shared.",
        "",
        "RAG: inspector cases are real traces in `results/v4_improved/rag_case_studies.json`.",
        "",
        "Repository: 12 multi-file items × 3 models are scored in `results/v4_improved/repository_evaluation.json`. Live path is App.jsx → /api/chat → process_question → retrieve_documents → build_context → generate_response. `backend/rag/` is legacy.",
        "",
        "Limitations: rubric scoring is automatic; GPU is not measured; local hardware affects latency; 28 tasks are not a universal workload.",
        "",
        "Reproduce: `python week4_evaluation/scripts/run_pipeline.py` then open `week4_evaluation/dashboard/index.html`.",
        "",
        "Artifacts: `results/raw/week4_v4_improved.json`, `exports/week4_evaluation.xlsx`, `reports/week4_v4_report.md`, `reports/week4_v4_summary.md`, `dashboard/index.html`.",
    ]
    (REPORTS_DIR / "week4_v4_summary.md").write_text("\n".join(summary_lines), encoding="utf-8")
    return report_path


def write_dashboard_data(summary, rag, repo, resources):
    DASHBOARD_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "authoritative_version": EVALUATION_VERSION,
        "historical_versions_excluded": [
            "week4_v3_controlled (historical; not overwritten)",
            "archive_pre_v3 (V1/V2 raw scores — not used in this dashboard)",
        ],
        "summary": summary,
        "rag_cases": rag.get("cases", []),
        "repository": {
            "aggregates": repo.get("aggregates", {}),
            "call_graph": repo.get("actual_call_graph", []),
            "per_question": [
                {
                    "question_id": item.get("question_id"),
                    "model": item.get("model"),
                    "question": item.get("question"),
                    "expected_files": item.get("expected_files"),
                    "retrieved_files": item.get("retrieved_files"),
                    "identified_files_in_answer": item.get("identified_files_in_answer"),
                    "file_identification_score": item.get("file_identification_score"),
                    "dependency_reasoning_score": item.get("dependency_reasoning_score"),
                    "overall_score": item.get("overall_score"),
                    "model_answer": item.get("model_answer"),
                }
                for item in (repo.get("per_question") or [])
            ],
        },
        "v3_v4_comparison": summary.get("v3_v4_comparison"),
        "gpu_note": resources.get("gpu_note"),
        "model_display": MODEL_DISPLAY,
        "interpretations": category_interpretation(summary["category_matrix"]),
        "methodology": {
            "embedding_model": EMBEDDING_MODEL,
            "top_k": TOP_K,
            "temperature": TEMPERATURE,
            "num_predict": NUM_PREDICT,
            "shared_retrieval": True,
        },
    }
    json_path = DASHBOARD_DATA_FILE
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    js_path = DASHBOARD_DIR / "benchmark_data.js"
    js_path.write_text(
        "window.BENCHMARK = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    return js_path, payload


def main():
    dataset = load_dataset()
    raw = load_json(RAW_RESULTS_FILE)
    quality = load_json(QUALITY_SCORES_FILE)
    retrieval = load_json(RETRIEVAL_METRICS_FILE)
    resources = load_json(RESOURCE_METRICS_FILE)
    rag = load_json(RAG_CASE_STUDIES_FILE)
    repo = load_json(REPOSITORY_EVAL_FILE)
    summary = build_summary(raw, quality, retrieval, resources, repo, quality["aggregates"].get("code_generation"))
    comparison = build_v3_v4_comparison(summary)
    summary["v3_v4_comparison"] = comparison
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    RAW_CANONICAL_COPY.parent.mkdir(parents=True, exist_ok=True)
    RAW_CANONICAL_COPY.write_text(
        json.dumps(raw, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    SUMMARY_FILE.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    V3_V4_COMPARISON_FILE.parent.mkdir(parents=True, exist_ok=True)
    V3_V4_COMPARISON_FILE.write_text(
        json.dumps(comparison, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    excel = write_excel(dataset, raw, quality, retrieval, resources, rag, repo, summary)
    report = write_reports(summary, quality, retrieval, resources, rag, repo)
    data_path, _payload = write_dashboard_data(summary, rag, repo, resources)
    print(f"Wrote {SUMMARY_FILE}")
    print(f"Wrote {excel}")
    print(f"Wrote {report}")
    print(f"Wrote {data_path}")
    print("Dashboard HTML is week4_evaluation/dashboard/index.html")


if __name__ == "__main__":
    main()
